import json, copy, os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from kich_ban import SCRIPT
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

D = json.load(open('/tmp/claude-0/pp/slides.json'))
PX = 9525
E = lambda v: Emu(int(round(v * PX)))
rgb = lambda h: RGBColor.from_string(h)
PART = ['1 · Dân chủ XHCN ở Việt Nam', '2 · Nhà nước pháp quyền XHCN Việt Nam', '3 · Phát huy & xây dựng hiện nay']
SANS, SERIF = 'Arial', 'Georgia'
SC = [1.0]

prs = Presentation()
prs.slide_width, prs.slide_height = E(1920), E(1080)
blank = prs.slide_layouts[6]

def shape(sl, kind, x, y, w, h, fill=None, line=None, lw=1.5, dash=None, adj=None):
    s = sl.shapes.add_shape(kind, E(x), E(y), E(w), E(h))
    s.shadow.inherit = False
    if fill: s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    else: s.fill.background()
    if line:
        s.line.color.rgb = rgb(line); s.line.width = Pt(lw * .75)
        if dash: s.line.dash_style = dash
    else: s.line.fill.background()
    if adj is not None:
        for i, a in enumerate(adj): s.adjustments[i] = a
    return s

def write(s, runs, size, color, font=SANS, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, ml=22, mt=14, italic=False, bold=False):
    tf = s.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = E(ml); tf.margin_top = tf.margin_bottom = E(mt)
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]; p.alignment = align
    for r in runs:
        if r.get('nl'):
            p = tf.add_paragraph(); p.alignment = align; continue
        run = p.add_run(); run.text = r['t']
        f = run.font; f.name = font; f.size = Pt(r.get('size', size) * .75 * SC[0])
        f.bold = bool(r.get('b', bold)); f.italic = italic
        f.color.rgb = rgb(r.get('c', color if not r.get('r') else '8C2F2B'))

def dash_of(name):
    from pptx.enum.dml import MSO_LINE
    return {'dash': MSO_LINE.DASH, 'dot': MSO_LINE.ROUND_DOT}.get(name)

def line_ends(conn, head, tail):
    ln = conn.line._get_or_add_ln()
    for tag, on in (('a:headEnd', head), ('a:tailEnd', tail)):
        if on:
            e = etree.SubElement(ln, qn(tag)); e.set('type', 'triangle'); e.set('w', 'med'); e.set('len', 'med')

def add_timing(slide, clicks):
    """clicks: list of (entr_ids, exit_ids)"""
    P = 'http://schemas.openxmlformats.org/presentationml/2006/main'
    n = [2]
    def nid(): n[0] += 1; return n[0]
    def eff(spid, first, exit_):
        c1, c2, c3 = nid(), nid(), nid()
        node = 'clickEffect' if first else 'withEffect'
        if exit_:
            return (f'<p:par><p:cTn id="{c1}" presetID="1" presetClass="exit" presetSubtype="0" fill="hold" nodeType="{node}"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
                    f'<p:set><p:cBhvr><p:cTn id="{c2}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="hidden"/></p:to></p:set>'
                    f'</p:childTnLst></p:cTn></p:par>')
        c4 = nid()
        return (f'<p:par><p:cTn id="{c1}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" nodeType="{node}"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
                f'<p:set><p:cBhvr><p:cTn id="{c2}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
                f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{c3}" dur="300"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>'
                f'</p:childTnLst></p:cTn></p:par>')
    body = ''
    for ents, exits in clicks:
        a, b = nid(), nid()
        effs = ''
        first = True
        for sid in ents: effs += eff(sid, first, False); first = False
        for sid in exits: effs += eff(sid, first, True); first = False
        body += (f'<p:par><p:cTn id="{a}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>'
                 f'<p:par><p:cTn id="{b}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>{effs}</p:childTnLst></p:cTn></p:par>'
                 f'</p:childTnLst></p:cTn></p:par>')
    xml = (f'<p:timing xmlns:p="{P}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
           f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>{body}</p:childTnLst></p:cTn>'
           f'<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
           f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
           f'</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>')
    slide._element.append(etree.fromstring(xml))

SCRIPT_BY = {x['id']: x for x in SCRIPT}
PAGE = {}
for d in D:
    sl = prs.slides.add_slide(blank)
    sl.background.fill.solid(); sl.background.fill.fore_color.rgb = rgb('F4E9D2')
    # khung + header
    shape(sl, MSO_SHAPE.RECTANGLE, 14, 14, 1892, 1052, None, 'C9A24B', 2.5)
    if not d['cover']:
        shape(sl, MSO_SHAPE.RECTANGLE, 0, 0, 1920, 112, '8C2F2B')
        shape(sl, MSO_SHAPE.RECTANGLE, 0, 112, 1920, 5, 'C9A24B')
        t = shape(sl, MSO_SHAPE.RECTANGLE, 60, 8, 1800, 96)
        write(t, [{'t': d['title']}], 44, 'F6E3AA', SERIF, anchor=MSO_ANCHOR.MIDDLE, ml=20, mt=0, bold=True)
    else:
        t = sl.shapes.add_textbox(E(180), E(300), E(1560), E(380))
        write(t, [{'t': d['title']}], 76, '8C2F2B', SERIF, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 0, 0, bold=True)
        shape(sl, MSO_SHAPE.RECTANGLE, 780, 700, 360, 5, 'C9A24B')
    SC[0] = 1.0
    # thanh tiến độ
    for i, nm in enumerate(PART):
        on = d['part'] == 4 or d['part'] == i + 1
        done = d['part'] != 4 and d['part'] > i + 1
        fill = '8C2F2B' if on else ('E6CF93' if done else 'EADBBB')
        s = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, 70 + i * 530, 1003, 520, 44, fill, 'C9A24B' if on else 'B8AFA0', 1.5, adj=[.5])
        write(s, [{'t': nm}], 20, 'FFFFFF' if on else ('6A1F1C' if done else '6B6B6B'), SANS, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 4, 0, bold=True)
    # cạnh
    edge_ids = []   # (need, [shape ids])
    ST = {'arrow': ('6B6B6B', 2.5, None, False), 'conv': ('8C2F2B', 4, None, False),
          'cross': ('1F3A5F', 2.5, 'dash', False), 'vs': ('6B6B6B', 3, 'dot', True)}
    for e in d['edges']:
        if not e['on']: continue
        col, w, dash, both = ST[e['kind']]
        ids = []
        pts = e['pts']
        for i in range(len(pts) - 1):
            (x1, y1), (x2, y2) = pts[i], pts[i + 1]
            c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
            c.line.color.rgb = rgb(col); c.line.width = Pt(w * .75)
            if dash: c.line.dash_style = dash_of(dash)
            line_ends(c, both and i == 0, i == len(pts) - 2)
            ids.append(c.shape_id)
        if e['lab']:
            L = e['lab']
            lb = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, L['x'], L['y'], L['w'], 26, 'FBF6EC', '1F3A5F', 1.5, adj=[.5])
            write(lb, [{'t': L['t']}], 17, '1F3A5F', SANS, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 0, 0, bold=True)
            ids.append(lb.shape_id)
        edge_ids.append((e['need'], ids))
    # node
    SC[0] = .95
    node_ids = {}   # step -> ids ; placeholders
    placeholders = {}
    for n in d['nodes']:
        x, y, w, h, fs, t = n['x'], n['y'], n['w'], n['h'], n['fs'], n['type']
        ids = []
        if t == 'hub' and n.get('phOnly'):
            # trang tách (chưa tới đáp án): chỉ hiện ô dấu hỏi, không có hiệu ứng
            s = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, 'D9B9B0', 'C9A24B', 4, adj=[min(.25, 24 / min(w, h))])
            write(s, [{'t': '?'}], 64, '8C2F2B', SANS, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 0, 0, bold=True)
            continue
        if t == 'hub':
            s = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, '8C2F2B', 'C9A24B', 4, adj=[min(.25, 24 / min(w, h))])
            write(s, n['runs'], fs * 1.05, 'FFFFFF', SANS, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 28, 20, bold=True)
            for r in s.text_frame.paragraphs[0].runs: r.font.color.rgb = rgb('FFFFFF')
            if n['s'] > d['base'] and not __import__('os').environ.get('NOPH'):
                ph = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, 'D9B9B0', 'C9A24B', 4, adj=[min(.25, 24 / min(w, h))])
                write(ph, [{'t': '?'}], 64, '8C2F2B', SANS, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 0, 0, bold=True)
                placeholders.setdefault(n['s'], []).append(ph.shape_id)
        elif t == 'head':
            s = shape(sl, MSO_SHAPE.RECTANGLE, x, y, w, h, None, None)
            write(s, n['runs'], fs * 1.0, '8C2F2B', SANS, anchor=MSO_ANCHOR.BOTTOM, ml=4, mt=4, bold=True)
            shape(sl, MSO_SHAPE.RECTANGLE, x, y + h - 3, w, 3, 'C9A24B')
        elif t == 'quote':
            s = shape(sl, MSO_SHAPE.ROUNDED_RECTANGULAR_CALLOUT, x, y, w, h, 'FFF6DC', 'C9A24B', 2, adj=[-0.5 + 52 / w, 0.5 + 16 / h, .16])
            write(s, n['runs'], fs, '2B2622', SERIF, italic=True, ml=28, mt=18)
        elif t == 'observe':
            s = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, 'FFFFFF', '6B6B6B', 2, dash_of('dash'), [.5])
            write(s, n['runs'], fs * .9, '6B6B6B', SANS, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 12, 4)
        elif t == 'ghost':
            s = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, None, '1F3A5F', 2, dash_of('dash'), [.12])
            write(s, n['runs'], fs, '1F3A5F', SANS)
            if n.get('alertStep'):
                ov = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, 'FBE3DF', 'D33A2C', 3, dash_of('dash'), [.12])
                write(ov, n['runs'], fs, 'D33A2C', SANS, bold=True)
                for r in ov.text_frame.paragraphs[0].runs: r.font.color.rgb = rgb('D33A2C')
                node_ids.setdefault(n['alertStep'], []).append(ov.shape_id)
            for r in s.text_frame.paragraphs[0].runs:
                if r.font.color.rgb == rgb('8C2F2B'): r.font.color.rgb = rgb('1F3A5F')
        elif t == 'chips':
            s = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, 'FFFBF1', 'B8AFA0', 1.5, adj=[.06])
            runs = [{'t': n['chip']['head'], 'b': True, 'c': '8C2F2B', 'size': 22}, {'nl': 1}]
            for i, it in enumerate(n['chip']['items']):
                runs.append({'t': it + ('   ·   ' if i < len(n['chip']['items']) - 1 else ''), 'size': 17})
            write(s, runs, 17, '2B2622', ml=16, mt=10)
        elif t == 'xcard':
            s = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, 'FFFBF1', '1F3A5F', 1.5, adj=[.12])
            c = n['xc']
            write(s, [{'t': c['code'] + '  ', 'b': True, 'c': '1F3A5F', 'size': 22}, {'t': c['b'], 'b': True, 'size': 18}, {'nl': 1}, {'t': c['e'], 'c': '6B6B6B', 'size': 16}], 18, '2B2622', ml=14, mt=8, anchor=MSO_ANCHOR.MIDDLE)
        elif d['cover']:
            s = shape(sl, MSO_SHAPE.RECTANGLE, 260, y, 1400, h, None, None)
            write(s, n['runs'], 30, '6B6B6B', SANS, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        else:
            lc = 'D33A2C' if n['alert'] else 'B8AFA0'
            s = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, 'FFFBF1', lc, 2.5 if n['alert'] else 1.5, dash_of('dash') if n['alert'] else None, [min(.2, 12 / min(w, h))])
            write(s, n['runs'], fs, 'D33A2C' if n['alert'] else '2B2622')
        ids.append(s.shape_id)
        if n['who']:
            ww = len(n['who']) * 9.5 + 30
            wb = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, x + 24, y - 15, ww, 28, '8C2F2B', 'C9A24B', 2, adj=[.5])
            write(wb, [{'t': n['who']}], 17, 'FFFFFF', SANS, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 0, 0, bold=True)
            ids.append(wb.shape_id)
        cx = x + w + 10
        for code in reversed(n['xs']):
            bw = 44 + (12 if len(code) > 2 else 0)
            cx -= bw
            bg = shape(sl, MSO_SHAPE.ROUNDED_RECTANGLE, cx, y - 15, bw, 30, '1F3A5F', 'F4E9D2', 2, adj=[.5])
            write(bg, [{'t': code}], 16, 'FFFFFF', SANS, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 0, 0, bold=True)
            ids.append(bg.shape_id); cx -= 5
        node_ids.setdefault(n['s'], []).extend(ids)
    # hoạt cảnh: mỗi bước = 1 cú bấm
    steps = sorted({s for s in node_ids if s > d['base']} | {nd for nd, _ in edge_ids if nd > d['base']} | set(k for k in placeholders if k > d['base']))
    clicks = []
    for st in steps:
        ents = list(node_ids.get(st, []))
        for nd, ids in edge_ids:
            if nd == st: ents += ids
        clicks.append((ents, placeholders.get(st, [])))
    # các node / cạnh có step <= base hoặc cạnh có need < ... : giữ tĩnh
    # đặt các phần tử hiện ở bước <= base tĩnh (không có hiệu ứng)
    if clicks: add_timing(sl, clicks)
    # ghi chú người nói: lấy từ kịch bản (slide tách 2 trang thì chia lời theo "split")
    sc = SCRIPT_BY[d['rawid']]
    k = PAGE.get(d['rawid'], 0); PAGE[d['rawid']] = k + 1
    say = sc['say']
    if 'split' in sc: say = say[:sc['split']] if k == 0 else say[sc['split']:]
    note = [f"{sc['who']} · {sc['time']}", f"Thao tác: {sc['act']}", f"(Trang này: {len(clicks)} lần bấm)" if clicks else "(Trang này: không có hiệu ứng)", ''] + say
    if sc.get('next') and ('split' not in sc or k > 0): note += ['', 'Chuyển ý: ' + sc['next']]
    sl.notes_slide.notes_text_frame.text = '\n'.join(note)

prs.save(__import__('os').environ.get('OUT','/home/user/mln122/mln133/MLN131_Chuong4_MucIII.pptx'))
print('ok', len(D))
