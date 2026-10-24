# Tạo Kich_ban_thuyet_trinh_tung_slide.docx từ kich_ban.py
import os, sys
from docx import Document
from docx.shared import Pt, RGBColor, Cm
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from kich_ban import SCRIPT

RED, GREY = RGBColor(0x8C, 0x2F, 0x2B), RGBColor(0x6B, 0x6B, 0x6B)
OUT = os.environ.get('OUT', os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'Kich_ban_thuyet_trinh_tung_slide.docx'))

doc = Document()
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2.2); s.top_margin = s.bottom_margin = Cm(2)
st = doc.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(12)

def para(text, size=12, bold=False, italic=False, color=None, after=6):
    p = doc.add_paragraph(); r = p.add_run(text)
    r.font.size = Pt(size); r.bold = bold; r.italic = italic
    if color: r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(after)
    return p

para('KỊCH BẢN THUYẾT TRÌNH THEO TỪNG SLIDE', 18, True, color=RED)
para('MLN131 · Chương 4 · Mục III. Dân chủ xã hội chủ nghĩa và Nhà nước pháp quyền xã hội chủ nghĩa ở Việt Nam', 13, True)
para('Lớp: [điền] · Nhóm: [điền] · Thành viên: [Họ tên – MSSV] × 4 · Tổng thời lượng khoảng 32–33 phút', 11, color=GREY)
para('Cách dùng', 13, True, color=RED, after=2)
for t in ['Đây là lời nói gợi ý, không cần thuộc từng chữ. Nắm ý rồi nói bằng giọng của mình; chỗ nào thấy gượng thì cứ đổi cho tự nhiên.',
          'Những câu trong ngoặc kép là trích nguyên văn từ giáo trình, văn kiện hoặc Hồ Chí Minh — giữ đúng câu chữ.',
          '“(bấm)” là bấm → một lần trên pptx. Dòng “Thao tác” ghi tổng số lần bấm của slide đó.',
          'Lời nói này cũng đã được chép vào phần ghi chú (Notes) của từng slide trong pptx; bật Presenter View là thấy.',
          'Slide “Nước ta là nước dân chủ” và slide “Dân tự làm, hay giao cho người đại diện?” trong pptx được tách thành hai trang (trước và sau khi lớp trả lời).',
          'Thứ tự slide: Bìa → Nội dung trình bày → Phần 1 → Phần 2 (mở đầu bằng “Nhà nước pháp quyền là gì?”) → Phần 3 → Tổng kết → Hỏi đáp.']:
    p = para('• ' + t, 11, after=2)
doc.add_paragraph()

for i, sc in enumerate(SCRIPT, 1):
    para(f"Slide {i} — {sc['title']}", 14, True, color=RED, after=2)
    para(f"{sc['who']} · {sc['time']}", 11, True, color=GREY, after=2)
    para('Thao tác: ' + sc['act'], 11, italic=True, color=GREY, after=4)
    for t in sc['say']:
        para(t, 12, after=6)
    if sc.get('next'):
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(14)
        r = p.add_run('Chuyển ý: '); r.bold = True; r.font.color.rgb = RED; r.font.size = Pt(11)
        r = p.add_run('“' + sc['next'] + '”'); r.italic = True; r.font.size = Pt(11)
    else:
        doc.add_paragraph().paragraph_format.space_after = Pt(8)

doc.save(OUT)
print('ok', len(SCRIPT), OUT)
